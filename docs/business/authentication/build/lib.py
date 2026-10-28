import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY=RGBColor(0x1a,0x21,0x4d); BLUE=RGBColor(0x3b,0x7e,0xf6); YELLOW=RGBColor(0xfd,0xba,0x23)
LAV=RGBColor(0xed,0xf2,0xff); WHITE=RGBColor(0xff,0xff,0xff); GRAY=RGBColor(0x66,0x66,0x66)
TXT=NAVY

def font(run,size=11,bold=None,color=TXT,italic=None,name='Lato'):
    run.font.name=name
    rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),name)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color is not None: run.font.color.rgb=color

def spacing(p,before=0,after=4,line=None,keep=False):
    pf=p.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after)
    if line: pf.line_spacing=line
    if keep: pf.keep_with_next=True

def rich(p,text,size=11,color=TXT,base_bold=False,italic=False):
    """**bold** markup"""
    for i,part in enumerate(re.split(r'(\*\*[^*]+\*\*)',text)):
        if not part: continue
        b=part.startswith('**')
        r=p.add_run(part[2:-2] if b else part); font(r,size,bold=(b or base_bold),color=color,italic=italic)
    return p

def rule(doc,color='FDBA23',sz=1.0,after=6,keep=True):
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=keep; pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bt=OxmlElement('w:bottom')
    bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),str(int(sz*8))); bt.set(qn('w:space'),'1'); bt.set(qn('w:color'),color)
    b.append(bt); pPr.append(b); spacing(p,0,after); return p

def shade(cell,hexfill):
    tcPr=cell._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn('w:shd')): tcPr.remove(s)
    shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexfill); tcPr.append(shd)

def cell_margins(table,top=70,bottom=70,left=100,right=100):
    tblPr=table._tbl.tblPr; m=OxmlElement('w:tblCellMar')
    for k,v in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e=OxmlElement('w:'+k); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); m.append(e)
    tblPr.append(m)

def borders(table,color='EDF2FF',sz=4,val='single'):
    tblPr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for k in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+k); e.set(qn('w:val'),val); e.set(qn('w:sz'),str(sz)); e.set(qn('w:space'),'0'); e.set(qn('w:color'),color); b.append(e)
    tblPr.append(b)

def widths(table,ws):
    table.autofit=False
    tblPr=table._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    grid=table._tbl.tblGrid
    for i,gc in enumerate(grid.findall(qn('w:gridCol'))):
        if i<len(ws): gc.set(qn('w:w'),str(int(ws[i]*1440)))
    tw=tblPr.find(qn('w:tblW'))
    if tw is None: tw=OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'),str(int(sum(ws)*1440))); tw.set(qn('w:type'),'dxa')
    for row in table.rows:
        for i,w in enumerate(ws): row.cells[i].width=Inches(w)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)

def no_split(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:cantSplit'); e.set(qn('w:val'),'true'); trPr.append(e)

def set_cell(cell,text,size=10,color=TXT,bold=False,italic=False,align=None):
    cell.paragraphs[0].text=''
    parts=text.split('\n')
    for i,t in enumerate(parts):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        spacing(p,0,2 if i<len(parts)-1 else 0)
        if t.startswith('• '):
            p.paragraph_format.left_indent=Inches(0.12); p.paragraph_format.first_line_indent=Inches(-0.12)
        rich(p,t,size,color,base_bold=bold,italic=italic)
        if align: p.alignment=align
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

def table(doc,headers,rows,ws,size=10,first_col=True):
    t=doc.add_table(rows=1+len(rows),cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    borders(t); cell_margins(t); widths(t,ws)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'1A214D'); set_cell(c,h,size,WHITE,bold=True)
    repeat_header(t.rows[0])
    for c in t.rows[0].cells:
        for pp in c.paragraphs: pp.paragraph_format.keep_with_next=True
    for ri,r in enumerate(rows):
        fill='FFFFFF' if ri%2==0 else 'EDF2FF'
        no_split(t.rows[ri+1])
        for ci,v in enumerate(r):
            c=t.rows[ri+1].cells[ci]
            if ci==0 and first_col: shade(c,'EDF2FF'); set_cell(c,v,size,NAVY,bold=True)
            else: shade(c,fill); set_cell(c,v,size)
    after=doc.add_paragraph(); spacing(after,0,6)
    return t
