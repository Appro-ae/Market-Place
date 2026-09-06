#!/usr/bin/env python3
"""
build_gap_xlsx.py — build the Appro-branded gap-analysis workbook from a JSON spec.

  python3 build_gap_xlsx.py spec.json
  python3 build_gap_xlsx.py --example > spec.json

Spec:
{
  "file": "output/02-topic/20260903-002-key-cpo-gap.xlsx",
  "sheets": [
    {"name": "Gap Analysis", "title": "…", "subtitle": "…", "meta": "…",
     "headers": ["Ref","CPO Input Requirement","Description","Where AMP-2548 Covers It","Coverage","Gap & Feedback to BA","Priority"],
     "widths": [7, 26, 34, 40, 16, 60, 9],
     "rows": [["R1", "…", "…", "…", "Partially covered", "…", "P1"]],
     "status_col": 5, "bold_cols": [1], "row_height": 110, "legend": true}
  ]
}
Status values styled (brand RAG): Covered / Done → blue; Partially covered / Done — fix required /
Flagged → yellow; Not covered / Not Done → navy. No formulas are written (LibreOffice recalculation
is unreliable here) — put computed counts in the spec.
"""
import argparse
import json
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip3 install openpyxl")

NAVY, BLUE, YELLOW, LAVENDER, WHITE = "FF1A214D", "FF3B7EF6", "FFFDBA23", "FFEDF2FF", "FFFFFFFF"
HAIR = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
F = lambda size=10, bold=False, color=NAVY: Font(name="Lato", size=size, bold=bold, color=color)
FILL = lambda c: PatternFill("solid", start_color=c)
TOP = Alignment(vertical="top", wrap_text=True)
TOP_C = Alignment(vertical="top", horizontal="center", wrap_text=True)
MID_C = Alignment(vertical="center", horizontal="center", wrap_text=True)

STATUS = {
    "covered": (FILL(BLUE), F(bold=True, color=WHITE)),
    "done": (FILL(BLUE), F(bold=True, color=WHITE)),
    "partially covered": (FILL(YELLOW), F(bold=True)),
    "done — fix required": (FILL(YELLOW), F(bold=True)),
    "done - fix required": (FILL(YELLOW), F(bold=True)),
    "done (flagged)": (FILL(YELLOW), F(bold=True)),
    "flagged": (FILL(YELLOW), F(bold=True)),
    "not covered": (FILL(NAVY), F(bold=True, color=WHITE)),
    "not done": (FILL(NAVY), F(bold=True, color=WHITE)),
}

EXAMPLE = {
    "file": "gap-example.xlsx",
    "sheets": [{
        "name": "Gap Analysis",
        "title": "<Feature> — CPO input requirements vs <KEY>",
        "subtitle": "Gap analysis · prepared by PO · <date>",
        "meta": "Sources: <KEY> v<N>, baselines <keys>, precedents <keys>. Coverage: 0 Covered · 3 Partially · 3 Not covered",
        "headers": ["Ref", "CPO Input Requirement", "Description", "Where <KEY> Covers It", "Coverage", "Gap & Feedback to BA", "Priority"],
        "widths": [7, 26, 34, 40, 16, 60, 9],
        "rows": [["R1", "Staff sees own applications only", "…", "AC4 pre-filter …", "Partially covered", "No capture AC; no user attribute …", "P1"]],
        "status_col": 5, "bold_cols": [1], "row_height": 110, "legend": True
    }]
}


def build_sheet(ws, s):
    ncols = len(s["headers"])
    ws.cell(row=1, column=1, value=s.get("title", ws.title)).font = F(16, True)
    ws.cell(row=2, column=1, value=s.get("subtitle", "")).font = F(11, True, BLUE)
    ws.cell(row=3, column=1, value=s.get("meta", "")).font = F(9)
    for c in range(1, ncols + 1):
        ws.cell(row=4, column=c).fill = FILL(YELLOW)
    ws.row_dimensions[4].height = 4
    hr = 6
    for i, (h, w) in enumerate(zip(s["headers"], s.get("widths", [20] * ncols)), start=1):
        cell = ws.cell(row=hr, column=i, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = F(bold=True, color=WHITE), FILL(NAVY), MID_C, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 30
    status_col = s.get("status_col")
    bold_cols = set(s.get("bold_cols", [1]))
    r = hr + 1
    for ri, row in enumerate(s["rows"]):
        band = FILL(WHITE) if ri % 2 == 0 else FILL(LAVENDER)
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.border, cell.alignment, cell.fill = BORDER, TOP, band
            cell.font = F(bold=True) if ci in bold_cols else F()
            if status_col and ci == status_col and isinstance(val, str) and val.strip().lower() in STATUS:
                cell.fill, cell.font = STATUS[val.strip().lower()]
                cell.alignment = TOP_C
        ws.row_dimensions[r].height = s.get("row_height", 90)
        r += 1
    if s.get("legend"):
        r += 1
        ws.cell(row=r, column=1, value="Legend").font = F(bold=True)
        for j, (label, key) in enumerate([("Covered / Done", "covered"), ("Partially covered / fix required", "partially covered"),
                                          ("Not covered / Not Done", "not covered")]):
            c = ws.cell(row=r, column=2 + j, value=label)
            c.fill, c.font = STATUS[key]
            c.alignment, c.border = MID_C, BORDER
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.sheet_view.showGridLines = False


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", nargs="?")
    ap.add_argument("--example", action="store_true")
    args = ap.parse_args()
    if args.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return
    if not args.spec:
        ap.error("spec.json is required (or --example)")
    spec = json.load(open(args.spec, encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for s in spec["sheets"]:
        build_sheet(wb.create_sheet(s["name"][:31]), s)
    wb.save(spec["file"])
    wb2 = load_workbook(spec["file"])
    print(f"wrote {spec['file']}: " + ", ".join(f"{ws.title} ({ws.max_row}x{ws.max_column})" for ws in wb2.worksheets))


if __name__ == "__main__":
    main()
